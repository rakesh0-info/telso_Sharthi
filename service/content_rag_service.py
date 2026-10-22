import os
import time
import logging
from sqlalchemy.orm import Session
from google import genai
from google.genai import types
from docx import Document as DocxDocument
from pypdf import PdfReader

from database_config import SessionLocal
from databaseModel.Content_document_model import ContentDocument, ContentDocumentChunk
from databaseModel.Content_model import Content
from enums.contentEnum import UploadStateEnum

logger = logging.getLogger(__name__)

EMBEDDING_MODEL = "gemini-embedding-001"
EMBEDDING_DIMENSION = 3072
BATCH_SIZE = 10


def process_content_document_background(doc_id: int):
    db: Session = SessionLocal()
    doc = None
    try:
        doc = db.query(ContentDocument).filter(ContentDocument.id == doc_id).first()
        if not doc:
            logger.error("Content document %s was not found", doc_id)
            return

        doc.upload_status = UploadStateEnum.NOT_START
        linked_content = (
            db.query(Content).filter(Content.id == doc.content_id).first()
            if doc.content_id is not None
            else None
        )
        if linked_content:
            linked_content.file_upload_state = UploadStateEnum.NOT_START
        db.commit()

        file_extension = os.path.splitext(doc.file_path)[1].lower()
        if file_extension == ".pdf":
            reader = PdfReader(doc.file_path)
            text_parts = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text_parts.append(extracted)
            text = "\n".join(text_parts)
        elif file_extension == ".docx":
            word_document = DocxDocument(doc.file_path)
            text_parts = [paragraph.text for paragraph in word_document.paragraphs if paragraph.text]
            for table in word_document.tables:
                for row in table.rows:
                    text_parts.append(" ".join(cell.text for cell in row.cells if cell.text))
            text = "\n".join(text_parts)
        elif file_extension == ".txt":
            with open(doc.file_path, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()
        else:
            raise ValueError(f"Unsupported document format: {file_extension or 'unknown'}")

        if not text.strip():
            raise ValueError("No extractable text was found in the uploaded document")

        words = text.split()
        chunks = []
        chunk_size, overlap = 500, 50
        for i in range(0, len(words), chunk_size - overlap):
            chunk = " ".join(words[i:i + chunk_size])
            if chunk:
                chunks.append(chunk)

        api_key = os.getenv("GEMINI_API_KEY")
        if chunks and not api_key:
            raise RuntimeError("GEMINI_API_KEY is required to generate document embeddings")
        client = genai.Client(api_key=api_key) if chunks else None
        new_chunk_rows = []
        for i in range(0, len(chunks), BATCH_SIZE):
            batch_texts = chunks[i:i + BATCH_SIZE]

            if client is None:
                raise RuntimeError("Gemini client was not initialized for non-empty document chunks")
            embedding_res = client.models.embed_content(
                model=EMBEDDING_MODEL,
                contents=batch_texts,
                config=types.EmbedContentConfig(
                    task_type="RETRIEVAL_DOCUMENT"
                ),
            )

            embeddings = embedding_res.embeddings or []
            if len(embeddings) != len(batch_texts):
                raise RuntimeError(
                    f"Gemini returned {len(embeddings)} embeddings for {len(batch_texts)} chunks"
                )

            for idx, (text_chunk, embedding) in enumerate(zip(batch_texts, embeddings)):
                values = embedding.values
                if values is None or len(values) != EMBEDDING_DIMENSION:
                    actual_dimension = len(values) if values is not None else 0
                    raise RuntimeError(
                        f"Gemini returned an embedding with dimension {actual_dimension}; "
                        f"expected {EMBEDDING_DIMENSION}"
                    )
                new_chunk_rows.append(
                    ContentDocumentChunk(
                        document_id=doc.id,
                        chunk_index=i + idx,
                        content=text_chunk,
                        embedding=values,
                    )
                )

            if i + BATCH_SIZE < len(chunks):
                time.sleep(1)

        db.query(ContentDocumentChunk).filter(
            ContentDocumentChunk.document_id == doc.id
        ).delete(synchronize_session=False)
        db.add_all(new_chunk_rows)
        doc.upload_status = UploadStateEnum.READY
        if linked_content:
            linked_content.file_upload_state = UploadStateEnum.READY
        db.commit()

    except Exception:
        db.rollback()
        logger.exception("Error processing content document %s", doc_id)
        if doc:
            try:
                failed_doc = db.query(ContentDocument).filter(ContentDocument.id == doc_id).first()
                if failed_doc:
                    failed_doc.upload_status = UploadStateEnum.FAILED
                    if failed_doc.content_id is not None:
                        linked_content = (
                            db.query(Content).filter(Content.id == failed_doc.content_id).first()
                        )
                        if linked_content:
                            linked_content.file_upload_state = UploadStateEnum.FAILED
                    db.commit()
            except Exception:
                db.rollback()
                logger.exception("Could not save failed status for content document %s", doc_id)
    finally:
        db.close()